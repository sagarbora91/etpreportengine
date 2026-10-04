"""Generate the synthetic Service Centre interim fixtures (Service interim S-2, decision 15, 3 Oct 2026).

  python scripts/service-centre/generate_fixtures.py

Input: scripts/service-centre/families.spec.json ONLY (frozen headers, canonical names and types). Nothing is read
from any real Service workbook, and every value is synthetic: jobs JOAW330SYN0001-0030, bills BIAW330SYN0001...,
customers "Sample Customer NN", phones "9XXXXXX0NN", e-mails "@example.invalid".

Output (byte-stable; re-running over an unchanged spec rewrites the same bytes):
  tests-dotnet/fixtures/service-interim/
    week1/Service Centre till 28 sep 2026/         S001-S040 (S038 header only) + 00_ control workbook
    week2/Service Centre till 05 oct 2026/         the same families, a week later (see WEEK2 below)
    undated/Service Centre/                        week1 S009 and S002 in a folder with no date (H1 failure)
    samedate-changed/Service Centre till 28 sep 2026/  week1 S009 with one row changed (IMPORT_PERIOD_ALREADY_PRESENT)
    gprc/Service Centre till 05 aug 2026/          S023 GPRC history, 1-5 Aug 2026 (lane L10, decision 16)
    gprc/Service Centre till 07 aug 2026/          raw GPRC CLAIM (S041), 3-7 Aug 2026, overlapping S023 on 3 and 5 Aug
  tests-dotnet/TestSupport/Service/ServiceFixtureExpectations.g.cs   Data rows per family per week

S041 GPRC CLAIM is a raw export only (the consolidated set never had it), so the week folders stay as they were and
S041 appears only under gprc/, in ETP's raw layout: sheet "GPRC Claims Report", header row 1, no Info sheet, and a
Transaction Date with a time of day.

Requires Python 3 and openpyxl. Every workbook is re-read after writing to assert its headers.
"""
import datetime as dt
import io
import json
import re
import sys
import zipfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
SPEC = Path(__file__).with_name("families.spec.json")
OUT = ROOT / "tests-dotnet" / "fixtures" / "service-interim"
EXPECTATIONS = ROOT / "tests-dotnet" / "TestSupport" / "Service" / "ServiceFixtureExpectations.g.cs"

WEEK1_DATE = dt.date(2026, 9, 28)
WEEK2_DATE = dt.date(2026, 10, 5)
WEEK1_FOLDER = "Service Centre till 28 sep 2026"
WEEK2_FOLDER = "Service Centre till 05 oct 2026"
COVERAGE_FROM = dt.date(2024, 4, 1)
FIXED_STAMP = (2026, 9, 29, 0, 0, 0)
FIXED_TIME = dt.datetime(*FIXED_STAMP)

SNAPSHOT_HISTORY = {"S006": "CLOSING STOCK", "S009": "PENDING REPORT", "S010": "PENDING DELIVERY"}
JOB_FIELDS = {"jobordernumber", "job_order_no", "joborder_number", "jonumber", "job_order_number", "jo_number"}
JOB_DATE_FIELDS = {"jodate", "joborder_date", "bookingdate", "booking_date"}
# The business-date column of each DateLog family (ServiceInterimFamilies.ReadRules).
BUSINESS_DATE = {
    "S003": "trans_date", "S004": "billingdate", "S007": "grn_date", "S008": "grn_date", "S013": "stm_date",
    "S019": "repairdate", "S022": "invoice_date", "S023": "transdate", "S024": "transdate", "S025": "transdate",
    "S026": "transdate", "S029": "repair_date", "S039": "transaction_date", "S040": "transaction_date",
    "S041": "transaction_date",
}
# Families with no consolidated workbook: never written to the week folders (S041 GPRC CLAIM, lane L10).
RAW_ONLY = {"S041"}
GPRC_HISTORY_FOLDER = "Service Centre till 05 aug 2026"
GPRC_CLAIM_FOLDER = "Service Centre till 07 aug 2026"
GPRC_CLAIM_FILE = "GPRC CLAIM 01.08.2026 TO 07.08.2026.xlsx"
GPRC_CLAIM_SHEET = "GPRC Claims Report"
STATUS_LABEL = {
    "S014": "DC", "S015": "IR", "S016": "RA", "S017": "RWR", "S018": "DELIVERED",
    "S031": "PD", "S032": "PR", "S033": "SRN", "S034": "REPAIRED", "S035": "SRNINV",
}
TENDERS = ["cashamount", "cardamount", "upi"]


def d(month, day):
    return dt.date(2026, month, day)


def job_no(n):
    return f"JOAW330SYN{n:04d}"


def job_date(n):
    return dt.date(2026, 8, 25) + dt.timedelta(days=n)


# ---------------------------------------------------------------------------------------------------------------
# Row plans. A row is a dict of overrides: "job" (job index), "date" (business or status date), "spare", "amount",
# "tender"; everything else is filled from the column name and type by fill().
# ---------------------------------------------------------------------------------------------------------------

def job_rows(jobs, days_after=3):
    return [{"job": n, "date": job_date(n) + dt.timedelta(days=days_after)} for n in jobs]


def dated_rows(dates, jobs=None):
    jobs = jobs or [None] * len(dates)
    return [{"date": day, "job": job} for day, job in zip(dates, jobs)]


def week1_plan():
    plan = {
        "S001": job_rows([1, 10, 12, 18]),
        "S002": job_rows([1, 2, 3, 4, 5, 6], days_after=0),
        "S003": [
            {"date": d(9, 26), "job": 18, "amount": 500},
            {"date": d(9, 27), "job": 19, "amount": 750},
            {"date": d(9, 27), "job": 19, "amount": 120},
            {"date": d(9, 28), "job": 20, "amount": 300},
        ],
        "S004": [{"date": day, "tender": tender, "amount": 1000 * (i + 1) + 100 * t}
                 for i, day in enumerate([d(9, 26), d(9, 27), d(9, 28)]) for t, tender in enumerate(TENDERS)],
        "S005": [{"date": d(9, 28)} for _ in range(3)],
        "S006": [{"date": WEEK1_DATE} for _ in range(4)],
        "S007": dated_rows([d(9, 24), d(9, 25), d(9, 27)]),
        "S008": dated_rows([d(9, 24), d(9, 25), d(9, 27)]),
        "S009": job_rows([1, 2, 3, 4, 5]),
        "S010": job_rows([6, 7, 8, 9]),
        "S011": job_rows([21, 22, 25]),
        "S012": job_rows([21, 22, 25]),
        "S013": dated_rows([d(9, 25), d(9, 26), d(9, 27)]),
        "S014": job_rows([10, 11, 29]),
        "S015": job_rows([12, 13, 3, 8]),
        "S016": job_rows([14, 15, 30]),
        "S017": job_rows([16, 17, 20]),
        # Job 18 has three line rows with different spare values.
        "S018": [dict(r, spare=s) for r, s in zip(job_rows([18, 18, 18]), [100, 250, 400])] + job_rows([19, 20]),
        "S019": dated_rows([d(9, 25), d(9, 26), d(9, 27)], [18, 19, 20]),
        "S020": job_rows([14, 15, 29]),
        "S021": job_rows([15, 30, 29]),
        "S022": dated_rows([d(9, 24), d(9, 25), d(9, 26)], [23, 24, 16]),
        "S023": dated_rows([d(9, 22), d(9, 23), d(9, 24)], [10, 11, 14]),
        "S024": dated_rows([d(9, 22), d(9, 23), d(9, 24)], [10, 11, 14]),
        "S025": dated_rows([d(9, 22), d(9, 23), d(9, 24)], [15, 17, 19]),
        "S026": dated_rows([d(9, 22), d(9, 23), d(9, 24)], [15, 17, 19]),
        "S027": [{"date": WEEK1_DATE} for _ in range(3)],
        "S028": [{"date": WEEK1_DATE} for _ in range(3)],
        "S029": dated_rows([d(9, 23), d(9, 24), d(9, 25)], [16, 17, 18]),
        "S030": job_rows([23, 23, 24]),
        "S031": job_rows([6, 7, 8]),
        "S032": job_rows([1, 2, 4]),
        "S033": job_rows([21, 22, 25]),
        "S034": job_rows([23, 24, 16]),
        "S035": job_rows([25, 21, 22]),
        "S036": job_rows([18, 19, 20], days_after=5),
        "S037": job_rows([23, 24, 16], days_after=5),
        "S038": [],
        "S039": dated_rows([d(9, 21), d(9, 22), d(9, 23)], [10, 12, 13]),
        "S040": dated_rows([d(9, 21), d(9, 22), d(9, 23)], [14, 15, 16]),
    }
    for index, row in enumerate(plan["S030"]):
        row["date"] = row["date"] + dt.timedelta(days=index)  # two running-test runs of job 23
    return plan


def week2_plan():
    """week1 plus exactly these changes:
    - jobs 1 and 2 leave S009 (pending repair) and appear in S018 DELIVERED;
    - job 12 leaves S015 (IR) and appears in S017 (RWR);
    - jobs 26, 27 and 28 enter S009;
    - S004: the 27 Sep card amount is restated (a changed money row);
    - S004: three rows are added for 5 Oct (cash 400, card 450, upi 125 = 975). The raw-CSV lane (L9) restates
      5 Oct with different money (raw 250 cash + 275 card = 525), so these amounts must stay different from those;
    - S003: one row is added (3 Oct).
    """
    plan = week1_plan()
    plan["S009"] = job_rows([3, 4, 5, 26, 27, 28])
    plan["S018"] = plan["S018"] + [{"job": n, "date": d(10, 2)} for n in (1, 2)]
    plan["S015"] = job_rows([13, 3, 8])
    plan["S017"] = job_rows([16, 17, 20]) + [{"job": 12, "date": d(10, 1)}]
    for row in plan["S004"]:
        if row["date"] == d(9, 27) and row["tender"] == "cardamount":
            row["amount"] += 50
    plan["S004"] = plan["S004"] + [{"date": d(10, 5), "tender": tender, "amount": amount}
                                   for tender, amount in zip(TENDERS, (400, 450, 125))]
    plan["S003"] = plan["S003"] + [{"date": d(10, 3), "job": 23, "amount": 410}]
    for code in ("S006", "S027", "S028"):
        for row in plan[code]:
            row["date"] = WEEK2_DATE
    return plan


# ---------------------------------------------------------------------------------------------------------------
# Cell values
# ---------------------------------------------------------------------------------------------------------------

def is_name(c):
    return bool(re.search(r"cust.*name|customername|endcustomername|^name$|technician_name", c))


def fill(code, column, row, index):
    c, kind = column["CanonicalField"], column["DataType"]
    if c in row.get("cells", {}):
        return row["cells"][c]
    job = row.get("job")
    nn = job if job is not None else index + 1
    day = row.get("date") or WEEK1_DATE

    if c in JOB_FIELDS:
        return job_no(job if job is not None else 90 + index)
    if c == "store_code":
        return "AW330"
    if kind == "Date":
        if c in ("sourceperiodfrom", "from_date"):
            value = dt.date(2026, 9, 1)
        elif c in ("sourceperiodto", "to_date"):
            value = day
        elif c == BUSINESS_DATE.get(code):
            value = day
        elif job is not None and c in JOB_DATE_FIELDS:
            value = job_date(job)
        elif job is not None and c == "edd":
            value = job_date(job) + dt.timedelta(days=10)
        elif job is not None and c in ("dop", "purchasedate"):
            value = dt.date(2025, 1, 1) + dt.timedelta(days=job)
        else:
            value = day
        if code == "S031":
            return value.strftime("%d-%m-%Y")  # S031 dates are text dd-MM-yyyy
        if code in ("S036", "S037") and c == "created_date":
            return f"{value.day}-{value.month}-{value.year}"  # text d-M-yyyy
        if c == BUSINESS_DATE.get(code) and row.get("time"):
            return dt.datetime.combine(value, row["time"])  # S041 Transaction Date keeps the time of day
        return dt.datetime(value.year, value.month, value.day)
    if kind == "Integer":
        return {"month": day.month, "year": day.year}.get(c, index + 1)
    if kind == "Decimal":
        if c in TENDERS:
            return row.get("amount", 0) if row.get("tender") == c else 0
        if code == "S004" and c != "totalamount":
            return 0  # the other tenders are empty, so TotalAmount equals the one tender
        if c in ("totalamount", "net_amount", "netamount", "netamount_incl_tax", "invoice_value", "amount", "total"):
            return row.get("amount", 100 + 10 * index)
        if c == "sparevalue":
            return row.get("spare", 100 + 10 * index)
        if c == "labourcharge":
            return 50
        if c in ("tax_amount", "taxamount", "total_tax", "totaltax"):
            return 18
        if c in ("quantity", "qty_shipped", "shipped_quantity", "received_quantity"):
            return 1 + index
        if c == "pendingnoofdays" and job is not None:
            return (WEEK1_DATE - job_date(job)).days
        return 10 + index
    if code == "S004" and c == "phonepe":
        return "0"  # PhonePe is a tender (typed Identifier in the spec), not a phone number
    if re.search(r"mobile|landline|contactnumber|(^|_)phone(_|$)|phone_?(no|number)", c):
        return f"9XXXXXX{nn:03d}"
    if re.search(r"e_?mail", c):
        return f"sample{nn:02d}@example.invalid"
    if "address" in c:
        return f"Sample Address {nn:02d}"
    if kind == "Identifier":
        if re.search(r"(billing|invoice|document|bill)_?(no|num|number)$|^billingnumber$|cashmemonumber", c):
            return f"BIAW330SYN{index + 1:04d}"
        return f"SYN-{code}-{index + 1:02d}"
    # Text
    if is_name(c):
        return f"Sample Customer {nn:02d}"
    if c in ("jostatus", "repairstatus", "current_status", "jo_status", "job_status", "joborder_status"):
        return STATUS_LABEL.get(code, "OPEN")
    if c in ("brand", "brandname", "brand_name"):
        return "Sample Brand " + "AB"[nn % 2]
    if c == "customerid" or c == "customer_id":
        return f"CUSTSYN{nn:03d}"
    if c == "sourcefile":
        return f"SAMPLE {code} 01.09.2026 TO {day:%d.%m.%Y}.xlsx"
    return "SAMPLE"


# ---------------------------------------------------------------------------------------------------------------
# Workbooks
# ---------------------------------------------------------------------------------------------------------------

def stable_zip(data):
    # openpyxl stamps each zip entry and docProps/core.xml with the current time; rewrite both.
    source, target = zipfile.ZipFile(io.BytesIO(data)), io.BytesIO()
    with zipfile.ZipFile(target, "w", zipfile.ZIP_DEFLATED) as output:
        for info in source.infolist():
            content = source.read(info.filename)
            if info.filename == "docProps/core.xml":
                content = re.sub(rb"(<dcterms:modified[^>]*>)[^<]*", rb"\g<1>2026-09-29T00:00:00Z", content)
                content = re.sub(rb"(<dcterms:created[^>]*>)[^<]*", rb"\g<1>2026-09-29T00:00:00Z", content)
            entry = zipfile.ZipInfo(info.filename, date_time=FIXED_STAMP)
            entry.compress_type = zipfile.ZIP_DEFLATED
            entry.external_attr = 0o600 << 16
            output.writestr(entry, content)
    return target.getvalue()


def new_workbook():
    import openpyxl
    workbook = openpyxl.Workbook()
    workbook.properties.creator = "openpyxl"
    workbook.properties.created = workbook.properties.modified = FIXED_TIME
    return workbook


def append_rows(sheet, rows):
    for row in rows:
        sheet.append(row)
        for cell in sheet[sheet.max_row]:
            if isinstance(cell.value, dt.datetime):
                has_time = cell.value.time() != dt.time()
                cell.number_format = "yyyy-mm-dd hh:mm:ss.000" if has_time else "yyyy-mm-dd"


def save(workbook, path, expected_sheets):
    import openpyxl
    buffer = io.BytesIO()
    workbook.save(buffer)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(stable_zip(buffer.getvalue()))
    check = openpyxl.load_workbook(path, read_only=True)
    for name, headers in expected_sheets.items():
        actual = list(next(check[name].iter_rows(max_row=1, values_only=True)))
        assert actual == headers, (path.name, name)
    check.close()


def family_workbook(family, rows, snapshot, history):
    code, headers = family["FamilyCode"], family["Headers"]
    workbook = new_workbook()
    data = workbook.active
    data.title = "Data"
    data.append(headers)
    values = [[fill(code, column, row, index) for column in family["Columns"]] for index, row in enumerate(rows)]
    append_rows(data, values)
    info = workbook.create_sheet("Info")
    for pair in (["Family ID", code], ["Coverage", f"{COVERAGE_FROM:%Y-%m-%d} to {snapshot:%Y-%m-%d}"],
                 ["Status", "Synthetic test fixture"]):
        info.append(pair)
    expected = {"Data": headers, "Info": ["Family ID", code]}
    if code in SNAPSHOT_HISTORY:
        sheet = workbook.create_sheet("Snapshot History")
        history_headers = headers + ["Snapshot_As_Of", "SourceFile"]
        sheet.append(history_headers)
        stem = SNAPSHOT_HISTORY[code]
        for as_of, as_of_rows in history:
            stamp = dt.datetime(as_of.year, as_of.month, as_of.day)
            block = [[fill(code, column, row, index) for column in family["Columns"]] + [stamp, f"{stem} {as_of:%d.%m.%Y}.csv"]
                     for index, row in enumerate(as_of_rows)]
            append_rows(sheet, block)
        expected["Snapshot History"] = history_headers
    return workbook, expected, len(values)


def control_workbook(families, counts, snapshot):
    workbook = new_workbook()
    sheet = workbook.active
    sheet.title = "Control"
    headers = ["Family ID", "File", "Rows", "Status"]
    sheet.append(headers)
    for family in families:
        code = family["FamilyCode"]
        sheet.append([code, f"{code}_{family['Name']}.xlsx", counts[code], "Synthetic"])
    info = workbook.create_sheet("Info")
    for pair in (["Family ID", "00"], ["Coverage", f"{COVERAGE_FROM:%Y-%m-%d} to {snapshot:%Y-%m-%d}"], ["Status", "Synthetic test fixture"]):
        info.append(pair)
    return workbook, {"Control": headers, "Info": ["Family ID", "00"]}


def load_spec():
    families = json.loads(SPEC.read_text(encoding="utf-8"))
    for family in families:
        code = family["FamilyCode"]
        assert re.fullmatch(r"S\d{3}", code) and code != "S038", code
        assert family["Headers"] == [column["SourceHeader"] for column in family["Columns"]], code
    # S038 SRNReport is a retired name whose header equals S011's; it is not in the spec, so it is built from S011.
    s011 = next(f for f in families if f["FamilyCode"] == "S011")
    s038 = dict(s011, FamilyCode="S038", ReportCode="S038", Name="SRNReport", TableName="etp_landing_s038")
    return sorted(families + [s038], key=lambda f: f["FamilyCode"])


def write_week(families, plan, folder, snapshot, histories):
    counts = {}
    families = [family for family in families if family["FamilyCode"] not in RAW_ONLY]
    for family in families:
        code = family["FamilyCode"]
        workbook, expected, count = family_workbook(family, plan[code], snapshot, histories.get(code, []))
        save(workbook, folder / f"{code}_{family['Name']}.xlsx", expected)
        counts[code] = count
    workbook, expected = control_workbook(families, counts, snapshot)
    save(workbook, folder / "00_Service_Centre_Consolidation_Control.xlsx", expected)
    return counts


def gprc_line(doc, day, job, value, ucp, time=None):
    """One GPRC claim line; the same cells fill S023 (technical headers) and S041 (readable headers)."""
    cells = {"documentnum": doc, "document_number": doc, "value": value, "ucpvalue": ucp, "ucp_value": ucp,
             "price": value, "netamountinctax": value, "net_amount_inc_tax": value, "accountnum": "GPRCCell",
             "account_number": "GPRCCell", "custname": "Sample Claim Party", "customer_name": "Sample Claim Party"}
    return {"date": day, "job": job, "time": time, "cells": cells}


def gprc_plans():
    """S023 history 1-5 Aug and S041 GPRC CLAIM 3-7 Aug 2026 (lane L10, decision 16 Q6/Q8).

    The union read takes S041 for every claim document it holds and S023 for the others:
    - GPAW330SYN0001 (1 Aug, 2 lines) is only in S023: read from S023;
    - GPAW330SYN0002 (3 Aug) is in both with the same line: read once, from S041;
    - GPAW330SYN0003 (5 Aug) is in both; S041 adds a second line: read from S041 (2 lines);
    - GPAW330SYN0004 (6 Aug, 1 line) and GPAW330SYN0005 (7 Aug, 2 lines) are only in S041.
    Union: 5 documents, 8 lines (2 from S023, 6 from S041); a plain UNION ALL would give 10 lines.
    """
    doc = lambda n: f"GPAW330SYN{n:04d}"
    history = [
        gprc_line(doc(1), d(8, 1), 10, 2, 150), gprc_line(doc(1), d(8, 1), 10, 10, 1570),
        gprc_line(doc(2), d(8, 3), 11, 2, 45),
        gprc_line(doc(3), d(8, 5), 14, 10, 450),
    ]
    claims = [
        gprc_line(doc(2), d(8, 3), 11, 2, 45, dt.time(11, 5, 12, 250000)),
        gprc_line(doc(3), d(8, 5), 14, 10, 450, dt.time(15, 35, 43, 146000)),
        gprc_line(doc(3), d(8, 5), 14, 2, 30, dt.time(15, 35, 43, 146000)),
        gprc_line(doc(4), d(8, 6), 12, 2, 25, dt.time(9, 0, 1, 5000)),
        gprc_line(doc(5), d(8, 7), 13, 10, 3465, dt.time(18, 20, 0)),
        gprc_line(doc(5), d(8, 7), 13, 2, 60, dt.time(18, 20, 0)),
    ]
    return history, claims


def raw_workbook(family, rows, sheet_name):
    """A workbook in ETP's raw export layout: one sheet, header row 1, no Info sheet."""
    code, headers = family["FamilyCode"], family["Headers"]
    workbook = new_workbook()
    data = workbook.active
    data.title = sheet_name
    data.append(headers)
    append_rows(data, [[fill(code, column, row, index) for column in family["Columns"]] for index, row in enumerate(rows)])
    return workbook, {sheet_name: headers}


def write_gprc(by_code):
    history, claims = gprc_plans()
    workbook, expected, _ = family_workbook(by_code["S023"], history, d(8, 5), [])
    save(workbook, OUT / "gprc" / GPRC_HISTORY_FOLDER / f"S023_{by_code['S023']['Name']}.xlsx", expected)
    workbook, expected = raw_workbook(by_code["S041"], claims, GPRC_CLAIM_SHEET)
    save(workbook, OUT / "gprc" / GPRC_CLAIM_FOLDER / GPRC_CLAIM_FILE, expected)


def write_expectations(week1, week2):
    lines = [
        "// <auto-generated>",
        "// Generated by scripts/service-centre/generate_fixtures.py from scripts/service-centre/families.spec.json.",
        "// Do not edit: change the generator and run it again.",
        "// </auto-generated>",
        "namespace Etp.Reporting.TestSupport.Service;",
        "",
        "internal static partial class ServiceFixtures",
        "{",
    ]
    for name, counts in (("Week1Rows", week1), ("Week2Rows", week2)):
        lines.append(f"    /// <summary>Data-sheet rows per family in the {name[:5].lower()} fixture folder.</summary>")
        lines.append(f"    public static IReadOnlyDictionary<string, int> {name} {{ get; }} = new Dictionary<string, int>(StringComparer.Ordinal)")
        lines.append("    {")
        lines += [f'        ["{code}"] = {count},' for code, count in counts.items()]
        lines.append("    };")
        lines.append("")
    lines[-1] = "}"
    EXPECTATIONS.parent.mkdir(parents=True, exist_ok=True)
    EXPECTATIONS.write_text("\r\n".join(lines) + "\r\n", encoding="utf-8", newline="")


def main():
    families = load_spec()
    one, two = week1_plan(), week2_plan()
    earlier = d(9, 21)
    histories1 = {code: [(earlier, one[code][:-1] or one[code]), (WEEK1_DATE, one[code])] for code in SNAPSHOT_HISTORY}
    histories2 = {code: histories1[code] + [(WEEK2_DATE, two[code])] for code in SNAPSHOT_HISTORY}

    week1 = write_week(families, one, OUT / "week1" / WEEK1_FOLDER, WEEK1_DATE, histories1)
    week2 = write_week(families, two, OUT / "week2" / WEEK2_FOLDER, WEEK2_DATE, histories2)

    by_code = {f["FamilyCode"]: f for f in families}
    for code in ("S009", "S002"):
        workbook, expected, _ = family_workbook(by_code[code], one[code], WEEK1_DATE, histories1.get(code, []))
        save(workbook, OUT / "undated" / "Service Centre" / f"{code}_{by_code[code]['Name']}.xlsx", expected)

    changed = [dict(row) for row in one["S009"]]
    changed[0]["date"] = changed[0]["date"] + dt.timedelta(days=1)  # one row changed: its dates move by a day
    workbook, expected, _ = family_workbook(by_code["S009"], changed, WEEK1_DATE, histories1["S009"])
    save(workbook, OUT / "samedate-changed" / WEEK1_FOLDER / "S009_PendingRepair.xlsx", expected)

    write_gprc(by_code)
    write_expectations(week1, week2)
    print(f"{OUT.relative_to(ROOT)}: {len(families) - len(RAW_ONLY)} families x 2 weeks + undated + samedate-changed + gprc; "
          f"{EXPECTATIONS.relative_to(ROOT)}")


if __name__ == "__main__":
    if len(sys.argv) != 1:
        sys.exit("usage: generate_fixtures.py")
    main()
