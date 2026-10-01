"""Generate the Service Centre (AW330) catalogue entries, migration 0038 and golden fixtures.

families.spec.json beside this script is the hand-reviewed source of truth. Every output is
deterministic and committed; re-running a sub-command over an unchanged spec rewrites the same bytes.

  python scripts/service-centre/generate.py catalogue   # append S families to EtpReportFamilies.json
  python scripts/service-centre/generate.py migration   # write database/migrations/0038_...sql
  python scripts/service-centre/generate.py fixtures    # write tests-dotnet/fixtures/etp-sample/Snnn_*.xlsx

Nothing is read from the real Service Centre folder. Fixture values are synthetic placeholders.
"""
import datetime
import io
import json
import re
import sys
import zipfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
SPEC = Path(__file__).with_name("families.spec.json")
CATALOGUE = ROOT / "src" / "Etp.Reporting.Import" / "Profiles" / "EtpReportFamilies.json"
MIGRATION = ROOT / "database" / "migrations" / "0038_service_centre_family_tables.sql"
FIXTURES = ROOT / "tests-dotnet" / "fixtures" / "etp-sample"

SQL_TYPES = {"Text": "nvarchar(max)", "Identifier": "nvarchar(max)", "Decimal": "decimal(19,4)", "Date": "date", "Integer": "int"}
KEY_ORDER = ["FamilyCode", "ReportCode", "Name", "IsTyped", "TableName", "PrimaryDateHeader", "BusinessUnit", "Headers", "Columns"]
# Families whose Data sheet carries a third consolidation sheet in the real folder.
SNAPSHOT_HISTORY = {"S006", "S009", "S010"}


def load_spec():
    families = json.loads(SPEC.read_text(encoding="utf-8"))
    for family in families:
        code = family["FamilyCode"]
        assert re.fullmatch(r"S\d{3}", code) and code != "S038", code
        assert family["ReportCode"] == code and family["TableName"] == f"etp_landing_{code.lower()}", code
        assert family["BusinessUnit"] == "SERVICE" and family["IsTyped"] is False, code
        assert family["Headers"] == [column["SourceHeader"] for column in family["Columns"]], code
        names = [column["CanonicalField"] for column in family["Columns"]]
        assert len(set(names)) == len(names), code
        assert all(re.fullmatch(r"[a-z0-9_]+", name) and "timestamp" not in name for name in names), code
        assert all(column["DataType"] in SQL_TYPES and column["IsRequired"] is False for column in family["Columns"]), code
    return families


def catalogue(families):
    text = CATALOGUE.read_text(encoding="utf-8")
    existing = json.loads(text)
    retail = [family for family in existing if not re.fullmatch(r"S\d{3}",family["FamilyCode"])]
    # The Retail entries must stay byte-identical: prove the serializer reproduces them first.
    retail_text = json.dumps(retail, indent=2, ensure_ascii=False) + "\n"
    assert text.startswith(retail_text[:-3]), "Retail catalogue entries would change"
    merged = retail + [{key: family[key] for key in KEY_ORDER} for family in families]
    CATALOGUE.write_text(json.dumps(merged, indent=2, ensure_ascii=False) + "\n", encoding="utf-8", newline="\n")
    print(f"{CATALOGUE.relative_to(ROOT)}: {len(retail)} Retail + {len(families)} Service families")


# Copied from 0018 (trigger) and 0025 (append procedure, DENY, GRANT); only the table and report code vary.
TRIGGER = ("EXEC(N'CREATE OR ALTER TRIGGER dbo.trg_{table}_locked ON dbo.[{table}] AFTER INSERT,UPDATE,DELETE AS BEGIN SET NOCOUNT ON; "
           "IF EXISTS(SELECT 1 FROM (SELECT import_file_id FROM inserted UNION SELECT import_file_id FROM deleted) x "
           "JOIN dbo.import_files f ON f.import_file_id=x.import_file_id JOIN dbo.daily_reporting_days d ON d.store_code=f.store_code "
           "AND d.business_date BETWEEN f.period_start AND f.period_end WHERE d.status=''LOCKED'') "
           "THROW 51243,''A day inside this export is finalised. Reopen it before changing its data.'',1; END');")
PROCEDURE_GUARDS = """AS BEGIN
 SET NOCOUNT ON; SET XACT_ABORT ON;
 IF COALESCE(IS_ROLEMEMBER(''etp_store_manager''),0)<>1 AND COALESCE(IS_ROLEMEMBER(''etp_owner''),0)<>1
    AND COALESCE(IS_SRVROLEMEMBER(''sysadmin''),0)<>1
  THROW 51420,''Owner or Store Manager permission is required.'',1;
 IF @@TRANCOUNT=0 THROW 51422,''Import writes require an enclosing transaction.'',1;
 IF NOT EXISTS(SELECT 1 FROM dbo.import_files f WITH(UPDLOCK,HOLDLOCK)
   JOIN dbo.import_batches b ON b.import_batch_id=f.import_batch_id
   WHERE f.import_file_id=@file AND f.is_superseded=0 AND f.data_truth_version=1 AND b.status=''Processing'')
  THROW 51422,''The source import is not open for writing.'',1;
 IF NOT EXISTS(SELECT 1 FROM dbo.source_lineage l JOIN dbo.import_files f ON f.import_file_id=l.import_file_id
   WHERE l.source_lineage_id=@lineage AND l.import_file_id=@file AND f.report_code=''{code}'')
  THROW 51422,''The source row does not belong to this report import.'',1;"""

HEADER = """-- Service Centre (AW330) landing tables for report families S001-S040 (S038 is retired and has no table).
-- Same shape as 0018 (table, file index, locked-day trigger) and 0025 (append procedure, DENY, GRANT EXECUTE).
-- Generated by scripts/service-centre/generate.py migration from scripts/service-centre/families.spec.json;
-- regenerate rather than edit. Guards make a re-run harmless; the migration runner supplies the transaction.
-- No dbo.stores row is seeded: an active AW330 would make scheduled Retail packs wait for a Retail sales export.
SET XACT_ABORT ON;
"""


def migration(families):
    blocks = [HEADER]
    for family in families:
        code, table, columns = family["FamilyCode"], family["TableName"], family["Columns"]
        names = [column["CanonicalField"] for column in columns]
        types = [SQL_TYPES[column["DataType"]] for column in columns]
        parameters = [f"@v{index}" for index in range(len(columns))]
        assert len(names) == len(types) == len(parameters) == len(family["Headers"]), code
        lines = [f"-- {code} {family['Name']} ({len(columns)} columns; "
                 + (f"period from {family['PrimaryDateHeader']})" if family["PrimaryDateHeader"] else "snapshot)"),
                 f"IF OBJECT_ID(N'dbo.{table}', N'U') IS NULL",
                 f"CREATE TABLE dbo.[{table}] (",
                 " etp_row_id bigint IDENTITY PRIMARY KEY,",
                 " import_file_id bigint NOT NULL REFERENCES dbo.import_files(import_file_id),",
                 " source_lineage_id bigint NOT NULL UNIQUE REFERENCES dbo.source_lineage(source_lineage_id),",
                 " content_key varchar(80) NOT NULL,"]
        lines += [f" [{name}] {sql_type} NULL" + ("," if index < len(names) - 1 else "") for index, (name, sql_type) in enumerate(zip(names, types))]
        lines += [");", "",
                  f"IF NOT EXISTS (SELECT 1 FROM sys.indexes WHERE name = N'IX_{table}_file' AND object_id = OBJECT_ID(N'dbo.{table}'))",
                  f" CREATE INDEX IX_{table}_file ON dbo.[{table}](import_file_id);", "",
                  TRIGGER.format(table=table), "",
                  f"EXEC(N'CREATE OR ALTER PROCEDURE dbo.append_{table}",
                  " @file bigint,@lineage bigint,@key varchar(80),"]
        lines += [f" {parameter} {sql_type}" + ("," if index < len(parameters) - 1 else "") for index, (parameter, sql_type) in enumerate(zip(parameters, types))]
        lines.append(PROCEDURE_GUARDS.format(code=code))
        lines.append(f" INSERT dbo.[{table}](import_file_id,source_lineage_id,content_key,{','.join(f'[{name}]' for name in names)})")
        lines.append(f" VALUES(@file,@lineage,@key,{','.join(parameters)});")
        lines += [" INSERT dbo.etp_import_content(import_file_id,source_row_number,content_key)",
                  " SELECT @file,source_row_number,@key FROM dbo.source_lineage WHERE source_lineage_id=@lineage;",
                  "END');", "",
                  f"DENY INSERT,UPDATE,DELETE ON dbo.[{table}] TO etp_store_manager,etp_viewer;",
                  f"GRANT EXECUTE ON dbo.append_{table} TO etp_store_manager,etp_owner;", ""]
        blocks.append("\n".join(lines))
    MIGRATION.write_text("\n".join(blocks), encoding="utf-8", newline="\n")
    print(f"{MIGRATION.relative_to(ROOT)}: {len(families)} tables, procedures and triggers")


SAMPLE_DATE = datetime.datetime(2026, 9, 28)
FIXED_TIME = datetime.datetime(2026, 9, 29, 0, 0, 0)


def sample_value(code, column):
    name, kind = column["CanonicalField"], column["DataType"]
    if kind == "Date":
        if code in ("S036", "S037") and name == "created_date": return "28-9-2026"
        if code == "S031": return "28-09-2026"
        return SAMPLE_DATE
    if kind == "Integer":
        return {"month": 9, "year": 2026}.get(name, 1)
    if kind == "Decimal":
        if re.search(r"inc_?l?_?tax|^totalamount$", name): return 118
        if name in ("tax_amount", "taxamount", "total_tax", "totaltax"): return 18
        return 100
    if kind == "Identifier":
        if name == "store_code": return "AW330"
        if re.search(r"job_?order|^jo_?(no|number)$|^jonumber$", name): return "JOAW330000000001"
        if re.search(r"(invoice|billing|document)_?(no|number)$", name): return "BIAW330000000001"
        if re.search(r"mobile|landline|(^|_)phone(_|$)", name): return "9XXXXXX000"
        return "SYNTH-0001"
    if re.search(r"cust.*name", name): return "Sample Customer"
    if re.search(r"e_?mail", name): return "sample@example.invalid"
    return "SAMPLE"


def stable_zip(data):
    # openpyxl stamps each zip entry and the core "modified" property with the current time;
    # rewrite both with a fixed stamp so regeneration is byte-stable.
    source, target = zipfile.ZipFile(io.BytesIO(data)), io.BytesIO()
    with zipfile.ZipFile(target, "w", zipfile.ZIP_DEFLATED) as output:
        for info in source.infolist():
            content = source.read(info.filename)
            if info.filename == "docProps/core.xml":
                content = re.sub(rb"(<dcterms:modified[^>]*>)[^<]*", rb"\g<1>2026-09-29T00:00:00Z", content)
            entry = zipfile.ZipInfo(info.filename, date_time=(2026, 9, 29, 0, 0, 0))
            entry.compress_type = zipfile.ZIP_DEFLATED
            output.writestr(entry, content)
    return target.getvalue()


def fixtures(families):
    import openpyxl
    for family in families:
        code = family["FamilyCode"]
        workbook = openpyxl.Workbook()
        workbook.properties.creator = "openpyxl"
        workbook.properties.created = workbook.properties.modified = FIXED_TIME
        data = workbook.active
        data.title = "Data"
        data.append(family["Headers"])
        data.append([sample_value(code, column) for column in family["Columns"]])
        for cell in data[2]:
            if isinstance(cell.value, datetime.datetime): cell.number_format = "yyyy-mm-dd"
        info = workbook.create_sheet("Info")
        for row in (["Family ID", code], ["Source", "AW330 sample 20260928"], ["Status", "Synthetic CI sample"]):
            info.append(row)
        if code in SNAPSHOT_HISTORY:
            workbook.create_sheet("Snapshot History").append(family["Headers"] + ["Snapshot_As_Of", "SourceFile"])
        buffer = io.BytesIO()
        workbook.save(buffer)
        path = FIXTURES / f"{code}_{family['Name']}.xlsx"
        path.write_bytes(stable_zip(buffer.getvalue()))
        check = openpyxl.load_workbook(path, read_only=True)
        headers = [value for value in next(check["Data"].iter_rows(max_row=1, values_only=True))]
        assert headers == family["Headers"], code
        check.close()
    print(f"{FIXTURES.relative_to(ROOT)}: {len(families)} Service fixtures")


if __name__ == "__main__":
    commands = {"catalogue": catalogue, "migration": migration, "fixtures": fixtures}
    if len(sys.argv) != 2 or sys.argv[1] not in commands:
        sys.exit("usage: generate.py catalogue|migration|fixtures")
    commands[sys.argv[1]](load_spec())
